# -*- coding: utf-8 -*-
"""test_greek_lease.py - round GL-1, 28 Oct 2026.

RELEASE 34 - THE GREEK LEASE WITHOUT ANYBODY ELSE'S DETAILS, AND
ALIVENTE'S DETAILS STORED ONCE. His decisions of 28 October:

    - The Greece template carried one past tenant's father's name, address,
      ΑΦΜ and ΔΟΥ, a floor, a parking space, a ΔΕΗ number, a two-year term
      and a third-year rent from 1.12.2023. All of it is a field now:
      the tenant's three on the tenant, the flat's three on the property
      (floors 1-3: P1/P2/P3, ΔΕΗ 3138572/3/4), the term worked out from
      the dates, the later rent a box in Generate - left out when blank.
    - Clause 44 no longer says the lease was signed by email.
    - Alivente's own details live once, on Administration -> System ->
      Company Details (superuser): registration HE283373 - Generate
      prefilled the placeholder HE123456 - and his passport, issued
      26-02-2022, expiring 25-02-2027, which the Greek lease names.
"""

# --- CONSOLE ENCODING ----------------------------------- 16 Sep 2026 --
# This file prints text it read out of the templates, and some of that
# text is not ASCII - projects/project_task_list.html carries a Greek
# heading behind the language switch, and it will not be the last. On
# Windows, Python writes stdout as cp1252 whenever it is not a UTF-8
# console, and cp1252 cannot encode Greek: the print itself raises
# UnicodeEncodeError and the run dies part-way through. A crash blocks a
# push exactly as hard as a failure and says far less about why.
#
# So keep the encoding the console really has - forcing UTF-8 only moves
# the problem to whoever decodes us - and change the ERROR HANDLER, so a
# character the console cannot draw arrives as a question mark instead of
# ending the run. stderr too, because a traceback is a print as well.
# Guarded, because stdout is not always a stream that can be told.
# See test_console_encoding.py.
import sys as _sys
for _stream in (_sys.stdout, _sys.stderr):
    try:
        _stream.reconfigure(errors='replace')
    except Exception:
        pass
# ------------------------------------------------------------------------

import os
import re
import sys

ROOT = os.getcwd()
sys.path.insert(0, ROOT)

SUFFIX = '.bak_greeklease'
ME = 'test_greek_lease.py'
PATCHER = 'apply_greek_lease.py'
PS1 = 'Push-PendingChanges.ps1'
MARK = 'GL-1, 28 Oct 2026'

J = os.path.join
TPL = J(ROOT, 'pages', 'templates')
MODELS = J(ROOT, 'pages', 'models.py')
FORMS = J(ROOT, 'pages', 'forms.py')
TENANTS = J(ROOT, 'pages', 'views', 'tenants.py')
PROPS = J(ROOT, 'pages', 'views', 'properties.py')
LEASE = J(ROOT, 'pages', 'views', 'lease_template.py')
URLS = J(ROOT, 'pages', 'urls.py')
VINIT = J(ROOT, 'pages', 'views', '__init__.py')
GEN = J(TPL, 'generate_lease_agreement.html')
TEDIT = J(TPL, 'tenant_edit.html')
PEDIT = J(TPL, 'properties_edit.html')
PADD = J(TPL, 'properties_add.html')
ADMIN = J(TPL, 'admin_apms.html')
GATE = J(ROOT, PS1)
EDITED = [MODELS, FORMS, TENANTS, PROPS, LEASE, URLS, VINIT, GEN, TEDIT,
          PEDIT, PADD, ADMIN]
MIGRATION = J(ROOT, 'pages', 'migrations', '0113_greek_lease_fields.py')
VIEW = J(ROOT, 'pages', 'views', 'company_details.py')
PAGE = J(TPL, 'company_details.html')
DOCX = J(TPL, 'lease_templates', 'greece_greek_lease_template.docx')
SHIPS = (MIGRATION, VIEW, PAGE)

passed = failed = skipped = 0


def ok(cond, msg, detail=''):
    global passed, failed
    if cond:
        passed += 1
        print('  ok   %s' % msg)
    else:
        failed += 1
        print('  FAIL %s' % msg)
        if detail:
            for line in str(detail).split('\n')[:8]:
                print('         %s' % line)
    return bool(cond)


def skip(msg, why):
    global skipped
    skipped += 1
    print('  --   %s  (%s)' % (msg, why))


def read(p):
    with open(p, encoding='utf-8-sig', errors='replace') as f:
        return f.read().replace('\r\n', '\n')


def was(p):
    b = p + SUFFIX
    if not os.path.isfile(b):
        raise IOError('%s is not on disk, so the control section cannot '
                      'run - and a control that does not run is the whole '
                      'reason the 17 Oct push failed.' % b)
    return read(b)


def head(t):
    print('')
    print('=' * 74)
    print(t)
    print('=' * 74)


print(__doc__.strip().splitlines()[0])

import alv_tree as T                                        # noqa: E402
from alv_rounds import as_left_by                          # noqa: E402

NOW = dict((p, as_left_by(p, SUFFIX, read)) for p in EDITED)
OLD = dict((p, was(p)) for p in EDITED)


def py(src):
    return T.python_code_only(src)


def claims(src):
    out = []
    m = py(src[MODELS])
    out.append((all('%s = models.CharField(max_length=' % f in m for f in (
        'tenant_father_name', 'tenant_afm', 'tenant_doy')),
        'the tenant keeps a father\'s name, ΑΦΜ and ΔΟΥ'))
    out.append((all('%s = models.CharField(max_length=' % f in m for f in (
        'prop_floor', 'prop_parking_space', 'prop_deh_number')),
        'the property keeps its floor, parking space and ΔΕΗ number'))
    out.append(('class CompanyDetails(models.Model):' in m
                and "'registration_number': 'HE283373'," in m
                and 'def current(cls):' in m,
                'Alivente\'s details are stored once, HE283373 by default'))
    f = py(src[FORMS])
    out.append(("'tenant_father_name', 'tenant_afm', 'tenant_doy'," in f
                and '"prop_deh_number",' in f,
                'the tenant and property forms keep the new fields'))
    out.append(('new_tenant.tenant_afm = original_tenant.tenant_afm' in
                py(src[TENANTS]),
                'Duplicate for Renewal copies them'))
    p = py(src[PROPS])
    out.append(("'prop_deh_number': property_obj.prop_deh_number or ''," in p
                and "'tenant_afm': tenant_obj.tenant_afm or ''," in p,
                'Generate is handed them'))
    lt = py(src[LEASE])
    out.append(('def _save_greek_back(' in lt and 'def greek_lease_terms(' in lt
                and "_save_greek_back(country, tenant_id, tenant_obj," in lt,
                'what Generate is told goes back, blanks only, and the term '
                'is worked out from the dates'))
    out.append(("'landlord_registration_number': additional_data.get("
                "'landlord_registration_number') or _co.registration_number,"
                in lt and "'landlord_passport_number':" in lt
                and "'company': CompanyDetails.current()," in lt,
                'the landlord comes from Company Details'))
    g = src[GEN]
    out.append(('id="gl-greek"' in g and 'function glGreek()' in g
                and "glGap(document.getElementById('gr-afm'))" in g
                and 'value="{{ company.registration_number }}"' in g
                and 'HE123456' not in g,
                'Generate asks for the Greek details and no longer prefills '
                'HE123456'))
    out.append(('id="gl-tenant-row"' in T.code_only(src[TEDIT])
                and '<input type="hidden" name="tenant_afm"' in src[TEDIT],
                'Tenant Edit shows them for a Greek property and carries '
                'them hidden elsewhere'))
    for pth, name in ((PEDIT, 'Property Edit'), (PADD, 'Property Add')):
        t = T.code_only(src[pth])
        out.append(('id="prop_floor"' in t and 'id="prop_parking_space"' in t
                    and 'id="prop_deh_number"' in t,
                    '%s has Floor, Parking Space and ΔΕΗ' % name))
    _sys = src[ADMIN][src[ADMIN].find('id="panel-system"'):]
    out.append(("{% url 'company_details' %}" in _sys
                and 'class="admin-btn admin-btn--reserved"' not in _sys,
                'Company Details takes the reserved System tile'))
    out.append(("name='company_details'" in py(src[URLS])
                and 'from .company_details import *' in py(src[VINIT]),
                '  and has a route'))
    return out


# ==========================================================================
head('1. THE SOURCE')
# ==========================================================================
for cond, msg in claims(NOW):
    ok(cond, msg)
for p in SHIPS:
    ok(os.path.isfile(p), '%s ships with the round' % os.path.relpath(p, ROOT))
mig = read(MIGRATION) if os.path.isfile(MIGRATION) else ''
ok("('pages', '0112_prop_owned_from')" in mig
   and mig.count('migrations.AddField(') == 6
   and "name='CompanyDetails'" in mig
   and 'migrations.RunPython(fill, migrations.RunPython.noop)' in mig,
   'migration 0113 adds six columns and the table after 0112, and fills them')
ok(all(s in mig for s in ("'prop_parking_space': 'P1'", "'3138572'",
                          "'3138573'", "'3138574'", "'HE283373'",
                          'datetime.date(2022, 2, 26)',
                          'datetime.date(2027, 2, 25)')),
   '  with his figures: P1-P3, ΔΕΗ 3138572-4, HE283373, the passport dates')
ok('if not getattr(p, k)' in mig and 'if not CD.objects.exists()' in mig,
   '  and only ever into a blank')
dated = [os.path.relpath(p, ROOT) for p in EDITED if MARK not in NOW[p]]
ok(not dated, 'the round dates itself in every file it edited',
   ', '.join(dated))
v = read(VIEW) if os.path.isfile(VIEW) else ''
ok('@user_passes_test(lambda u: u.is_superuser)' in v,
   'Company Details is superuser only')

# ==========================================================================
head('2. THE CONTROL - every claim above, asked of the backups')
# ==========================================================================
held = [m for c, m in claims(OLD) if c]
ok(not held, 'CONTROL: not one of this round\'s claims held before it',
   '\n'.join(held))

# ==========================================================================
head('3. THE GREEK TEMPLATE - nobody else\'s details left in it')
# ==========================================================================
try:
    from docx import Document
    have_docx = True
except Exception:                                             # noqa: BLE001
    have_docx = False
    skip('section 3', 'python-docx is not installed')


def doc_text(path):
    d = Document(path)
    return d, '\n'.join(p.text for p in d.paragraphs)


if have_docx:
    d, txt = doc_text(DOCX)
    gone = [s for s in ('Λαζάρου', '043454060', 'Χολαργού', '3138574',
                        '1.12.2023', 'χιλίων', 'ΑΤ4635239', '26-02-2022',
                        'μέσω ηλεκτρονικού ταχυδρομείου')
            if s in txt]
    ok(not gone, 'none of the past tenant\'s details, the old dates or '
                 '"signed by email" is left', ', '.join(gone))
    hl = [r.text for p in d.paragraphs for r in p.runs
          if r.font.highlight_color and r.text.strip()]
    ok(not hl, '  and nothing is highlighted for approval any more',
       ', '.join(hl[:6]))
    need = ('{{ tenant_father_name }}', '{{ tenant_afm }}', '{{ tenant_doy }}',
            '{{ property_floor_greek }}', '{{ property_parking_space }}',
            '{{ property_deh_number }}', '{{ lease_term_greek }}',
            '{{ lease_term_rent_greek }}', '{% if third_year_rent %}',
            '{{ landlord_passport_number }}', '{% if signed_electronically %}')
    miss = [s for s in need if s not in txt]
    ok(not miss, 'each is a field the lease fills', ', '.join(miss))
    ok('Αφαίας αρ. 15' in txt and '135 τ.μ.' in txt,
       '  and the building, which the three floors share, stays as written')
    bak = DOCX + SUFFIX
    if os.path.isfile(bak):
        _, old = doc_text(bak)
        ok('043454060' in old and 'Λαζάρου' in old,
           'CONTROL: the backup is the template that carried them')
    else:
        ok(False, 'CONTROL: the template has its backup', bak)

# ==========================================================================
head('4. DRIVEN - the lease, the save-back, the screens')
# ==========================================================================


class Rollback(Exception):
    """Nothing this suite writes survives it."""


have_db = False
try:
    import django
    from django.conf import settings as dj
    if not dj.configured:
        os.environ.setdefault('SECRET_KEY', 'test-only-not-a-secret')
        os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'mysite.settings')
        django.setup()
    from django.contrib.auth.models import User
    from django.db import connection, transaction
    from django.test import Client
    import pages.models as M
    from pages.views import lease_template as LT
    connection.ensure_connection()
    M.props.objects.exists()
    have_db = True
except Exception as _e:                                       # noqa: BLE001
    skip('section 4', 'no database: %s' % str(_e).split('\n')[0][:70])

if have_db and have_docx:
    import datetime as dt
    try:
        with transaction.atomic():
            co = M.CompanyDetails.current()
            ok(co.registration_number == 'HE283373'
               and co.passport_expires == dt.date(2027, 2, 25),
               'Company Details holds HE283373 and the passport dates')
            flat = M.props.objects.create(
                prop_name='GL1 Flat', prop_country='Greece',
                prop_available_for_rent='Yes', prop_address1='Afaias 15')
            ten = M.tenant.objects.create(
                prop=flat, tenant_type='Individual', tenant_name='GL1 Tenant',
                tenant_current='Yes', tenant_rental_type='Monthly',
                tenant_lease_start_date=dt.date(2026, 12, 1),
                tenant_lease_end_date=dt.date(2028, 11, 30), tenant_rent=950)
            ad = {'gr_father_name': 'Γεωργίου', 'gr_afm': '123456789',
                  'gr_doy': 'Ψυχικού', 'gr_floor': '2', 'gr_parking': 'P2',
                  'gr_deh': '3138573', 'third_year_rent': '1000',
                  'extension_period': '12', 'extension_date': '2029-11-30',
                  'start_date': '2026-12-01', 'end_date': '2028-11-30',
                  'monthly_rent': '950', 'security_deposit': '1900'}
            path = LT.generate_lease_document('greece', 'greek', flat, ten, ad)
            _, out = doc_text(path)
            ok('του Γεωργίου' in out and 'Α.Φ.Μ.: 123456789' in out
               and 'ΔΟΥ: Ψυχικού' in out,
               'the lease names THIS tenant\'s father, ΑΦΜ and ΔΟΥ')
            ok('του 2ου ορόφου' in out and 'με αριθμό P2 στον' in out
               and 'ΔΕΗ είναι 3138573.' in out,
               '  this flat\'s floor, parking space and ΔΕΗ number')
            ok('συμφωνείται για δύο (2) έτη.' in out
               and 'για τα πρώτα δύο (2) μισθωτικά έτη ήτοι' in out,
               '  a term worked out from the dates')
            ok('ενώ για το τρίτο μισθωτικό έτος ήτοι από την 1η Δεκεμβρίου '
               '2028' in out
               # GL-2, 28 Oct 2026 - his decision: a Greek lease writes
               # €1.000 with a dot. Either is this round's claim.
               and ('€1,000 (Χίλια Ευρώ)' in out
                    or '€1.000 (Χίλια Ευρώ)' in out),
               '  and the later rent, from the day after the term ends')
            ok('ΑΤ4635239 με ημερομηνία έκδοσης 26-02-2022 και ημερομηνία '
               'λήξης 25-02-2027' in out,
               '  and his passport from Company Details')
            ok('το υπέγραψαν.' in out and 'ταχυδρομείου' not in out,
               '  and clause 44 says signed, not signed by email')
            ad2 = dict(ad, third_year_rent='')
            _, out2 = doc_text(LT.generate_lease_document(
                'greece', 'greek', flat, ten, ad2))
            ok('τρίτο μισθωτικό έτος' not in out2
               and 'με τους ίδιους επακριβώς όρους.' in out2,
               'no later rent: neither clause mentions one')
            ok(LT.greek_lease_terms(dt.date(2026, 7, 15),
                                    dt.date(2027, 7, 14))[0] == 'ένα (1) έτος'
               and LT.greek_lease_terms(dt.date(2026, 1, 1),
                                        dt.date(2027, 6, 30))[0]
               == 'δεκαοκτώ (18) μήνες',
               'a one-year lease is "ένα (1) έτος", eighteen months is '
               '"δεκαοκτώ (18) μήνες"')

            got = LT._save_greek_back('greece', str(ten.pk), ten, flat, ad)
            ten.refresh_from_db()
            flat.refresh_from_db()
            ok(ten.tenant_afm == '123456789' and flat.prop_deh_number ==
               '3138573' and flat.prop_parking_space == 'P2',
               'what Generate was told goes onto the tenant and the flat')
            LT._save_greek_back('greece', str(ten.pk), ten, flat,
                                dict(ad, gr_afm='999', gr_deh='111'))
            ten.refresh_from_db()
            flat.refresh_from_db()
            ok(ten.tenant_afm == '123456789' and flat.prop_deh_number ==
               '3138573', '  blanks only - a value already there is never '
                          'overwritten')
            ok(LT._save_greek_back('cyprus', str(ten.pk), ten, flat, ad)
               is None, '  and a Cyprus lease writes nothing')

            u = User.objects.create_user('gl1', 'gl1@example.com', 'x',
                                         is_superuser=True, is_staff=True)
            c = Client(SERVER_NAME='localhost')
            c.force_login(u)
            body = c.get('/generate-lease-agreement/').content.decode()
            ok('value="HE283373"' in body and 'HE123456' not in body,
               'Generate prefills HE283373')
            r = c.post('/administration/company-details/', {
                'company_name': 'Alivente Limited',
                'registration_number': 'HE283373',
                'contact_person': 'Demetri Manias', 'phone': '+357',
                'email': 'x@example.com', 'address': 'Nicosia',
                'passport_number': 'ΑΤ0000001',
                'passport_issued': '2027-01-10',
                'passport_expires': '2037-01-09'})
            ok(r.status_code == 302 and M.CompanyDetails.current()
               .passport_number == 'ΑΤ0000001',
               'Company Details saves a renewed passport')
            _, out3 = doc_text(LT.generate_lease_document(
                'greece', 'greek', flat, ten, ad))
            ok('ΑΤ0000001 με ημερομηνία έκδοσης 10-01-2027' in out3,
               '  and the next lease names it')
            r = c.post('/administration/company-details/', {
                'company_name': 'A', 'passport_issued': '2027-01-10',
                'passport_expires': '2026-01-09'})
            ok(r.status_code == 200 and 'must come after' in
               r.content.decode(), '  an expiry before the issue is refused')
            plain = User.objects.create_user('gl1b', 'b@example.com', 'x')
            c2 = Client(SERVER_NAME='localhost')
            c2.force_login(plain)
            ok(c2.get('/administration/company-details/').status_code
               in (302, 403), 'and nobody but a superuser reaches it')
            raise Rollback()
    except Rollback:
        pass

# ==========================================================================
head('5. REGISTERED')
# ==========================================================================
import alv_rounds as R                                      # noqa: E402

ok(SUFFIX in R.ROUNDS, '%s is in alv_rounds.ROUNDS' % SUFFIX)
ok(ME in T.CONVERTED, '%s is in alv_tree.CONVERTED' % ME)
gate = read(GATE)
ok("'%s'" % ME in gate, '  and on the push gate suite list')
ok("Text = 'def _save_greek_back('" in gate,
   '  with a sentinel on the save-back')
wd = read(J(ROOT, 'test_waiting_down.py'))
nm = re.search(r'CONVERTED_N = (\d+)', wd)
ok(nm and int(nm.group(1)) == len(T.CONVERTED),
   '  and CONVERTED_N agrees with the register',
   '%s vs %d' % (nm.group(1) if nm else '?', len(T.CONVERTED)))
ok(os.path.isfile(J(ROOT, PATCHER)), 'the patcher is on disk')
nob = [os.path.relpath(p, ROOT) for p in EDITED + [DOCX]
       if not os.path.isfile(p + SUFFIX)]
ok(not nob, '  every file it edited has its backup', ', '.join(nob))

print('')
print('=' * 74)
print('%d passed, %d failed, %d skipped' % (passed, failed, skipped))
print('=' * 74)
sys.exit(1 if failed else 0)

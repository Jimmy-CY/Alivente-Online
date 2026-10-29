# -*- coding: utf-8 -*-
"""lease_generated.py - GE-1, 28 Oct 2026.

GENERATE -> E-SIGNING: THE TWO THINGS DONE TO A GENERATED LEASE.

  strip_signature_block   the e-signing version ends with one line saying it
                          was signed electronically, instead of the
                          hand-signature and witness lines - the signatures
                          are on the page ES-2 appends when everyone has
                          signed. His decision of 28 October.
  docx_to_pdf             LibreOffice, headless, on the server. His decision
                          of 28 October (option A): the PDF looks like the
                          Word document - heading, bold, footer, logo, page
                          numbers - where the HTML route lost all of them.

THE TEMPLATES ARE NOT EDITED FOR THIS. The block is found in the rendered
document and removed there, so the Word version every template produces is
byte-for-byte what it was, and the same four templates serve both buttons.

THE FONT. The templates ask for Trebuchet MS, which a server cannot carry.
The Dockerfile installs Inter and a fontconfig rule maps Trebuchet MS to it -
what the mock-ups he approved were rendered in. The Word download is
untouched and opens in Trebuchet on his PC.
"""
import os
import shutil
import subprocess
import tempfile

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

# Where the signature block starts, per template family. Matched on the
# LAST top-level paragraph that begins with one of these, case-folded.
ANCHORS = ('in witness whereof', 'οι συμβαλλομενοι', 'οι συμβαλλόμενοι')

ESIGN_LINE = {
    'english': ('Signed electronically. ',
                'The parties signed this agreement electronically on the '
                'dates shown on the signature page that follows, which '
                'records each signature, when and from where it was given, '
                'and the fingerprint of this document.'),
    'greek': ('Υπογράφηκε ηλεκτρονικά. ',
              'Οι συμβαλλόμενοι υπέγραψαν το παρόν ηλεκτρονικά, κατά τις '
              'ημερομηνίες που αναγράφονται στη σελίδα υπογραφών που '
              'ακολουθεί, η οποία καταγράφει κάθε υπογραφή, πότε και από '
              'πού δόθηκε, καθώς και το αποτύπωμα του παρόντος εγγράφου.'),
}

CONVERT_TIMEOUT = 90


class LeasePdfError(Exception):
    """A sentence a person can act on."""


def _text(el):
    return ''.join(t.text or '' for t in el.iter(W + 't'))


def _is_blank(el):
    if el.tag != W + 'p':
        return False
    return not _text(el).strip()


def strip_signature_block(docx_path, language):
    """Replace the hand-signature block with the e-signing line, in place.

    Raises ValueError when the block cannot be found - a template changed
    shape, and sending it with the old lines still in would be worse than
    stopping.
    """
    from docx import Document
    from docx.shared import Pt

    doc = Document(docx_path)
    body = doc.element.body
    els = [e for e in body if e.tag != W + 'sectPr']
    start = None
    for i, el in enumerate(els):
        if el.tag != W + 'p':
            continue
        t = ' '.join(_text(el).split()).lower()
        if t.startswith(ANCHORS):
            start = i
    if start is None:
        raise ValueError('The signature block was not found in this lease, '
                         'so it cannot be prepared for e-signing.')

    # The style to write in: the last real paragraph before the block.
    ref = None
    for el in reversed(els[:start]):
        if el.tag == W + 'p' and _text(el).strip():
            ref = el
            break

    for el in els[start:]:
        body.remove(el)
    # Trailing blank paragraphs, page breaks with them, would leave an
    # empty last page behind the line.
    while True:
        rest = [e for e in body if e.tag != W + 'sectPr']
        if rest and _is_blank(rest[-1]):
            body.remove(rest[-1])
        else:
            break

    font = size = None
    style = None
    if ref is not None:
        from docx.text.paragraph import Paragraph
        rp = Paragraph(ref, doc._body)
        style = rp.style
        for r in rp.runs:
            font = font or r.font.name
            size = size or r.font.size
    p = doc.add_paragraph(style=style)
    lead, rest_text = ESIGN_LINE.get(language, ESIGN_LINE['english'])
    for txt, bold in ((lead, True), (rest_text, False)):
        r = p.add_run(txt)
        r.bold = bold
        if font:
            r.font.name = font
        if size:
            r.font.size = size
    p.paragraph_format.space_before = Pt(14)
    doc.save(docx_path)
    return True


def soffice():
    for name in ('soffice', 'libreoffice'):
        found = shutil.which(name)
        if found:
            return found
    return None


def docx_to_pdf(docx_path, timeout=CONVERT_TIMEOUT):
    """The PDF bytes of a .docx, made by LibreOffice. Raises LeasePdfError."""
    exe = soffice()
    if not exe:
        raise LeasePdfError(
            'LibreOffice is not installed on this server, so the lease '
            'cannot be turned into a PDF here. Use Download Word, save it '
            'as a PDF, and attach it on the Send screen instead.')
    out = tempfile.mkdtemp(prefix='alv-pdf-')
    profile = tempfile.mkdtemp(prefix='alv-lo-')
    try:
        cmd = [exe, '--headless', '--norestore', '--nologo',
               '-env:UserInstallation=file://' + profile,
               '--convert-to', 'pdf', '--outdir', out, docx_path]
        try:
            subprocess.run(cmd, capture_output=True, timeout=timeout,
                           check=False)
        except subprocess.TimeoutExpired:
            raise LeasePdfError('Turning the lease into a PDF took too long '
                                'and was stopped. Try again in a minute.')
        base = os.path.splitext(os.path.basename(docx_path))[0] + '.pdf'
        pdf = os.path.join(out, base)
        if not os.path.isfile(pdf) or os.path.getsize(pdf) == 0:
            raise LeasePdfError('LibreOffice did not produce a PDF for this '
                                'lease. Use Download Word and attach a PDF '
                                'on the Send screen instead.')
        with open(pdf, 'rb') as fh:
            return fh.read()
    finally:
        shutil.rmtree(out, ignore_errors=True)
        shutil.rmtree(profile, ignore_errors=True)


def page_count(pdf_bytes):
    import pymupdf
    doc = pymupdf.open(stream=pdf_bytes, filetype='pdf')
    try:
        return doc.page_count
    finally:
        doc.close()


def first_page_png(pdf_bytes, dpi=40):
    """A small picture of page 1, for the Send screen's document card."""
    import pymupdf
    doc = pymupdf.open(stream=pdf_bytes, filetype='pdf')
    try:
        return doc[0].get_pixmap(dpi=dpi).tobytes('png')
    finally:
        doc.close()
